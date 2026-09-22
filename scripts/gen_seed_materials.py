# -*- coding: utf-8 -*-
"""Sinh mmv-warehouse/supabase/seed_materials.sql tu Material.xlsx.

Chay:  python scripts/gen_seed_materials.py
Can:   pip install openpyxl

Material.xlsx nam o git root va KHONG duoc commit (.gitignore loai *.xlsx),
nen phai co san file do tren may thi moi sinh lai duoc.
"""
import openpyxl, re, io, os, sys, unicodedata

# Chay tu bat ky dau: duong dan tinh theo vi tri file nay (scripts/ o git root).
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
XLSX = os.path.join(ROOT, "Material.xlsx")
OUT  = os.path.join(ROOT, "mmv-warehouse", "supabase", "seed_materials.sql")

CONSUMABLE = {
 'MMV1003021','MMV2009008','MMV1003020','MMV100090049','MMV09090057','MMV-OMO','MMV-VIM',
 'MMV10120080','MMV1111003','MMV10120066','MMV10040026','MMV10040025','MMV09090056',
 'MMV09090059','MMV1003017','MMV1003016','MMV10120082','MMV09090055','MMV09090060',
 'MMV110912','MMV10040030','STAPE09050042','STAPE09050042-1','RTAPE1209003','MMV1003012',
 'MMV10060040','ETAPE09050041','DSTAPE1209003','MMV10050033','MMV100100052','MMV1003011',
 'MMV-SIKA140',
}
ROLL = {'PF06060002','YUL-SP-3-008-1','CN-190321'}

def clean(v):
    if v is None:
        return None
    s = str(v).replace('\xa0', ' ')
    s = unicodedata.normalize('NFC', s)
    s = re.sub(r'\s+', ' ', s).strip()
    return s or None

def q(v):
    """Literal SQL: None -> null, con lai -> chuoi da escape."""
    if v is None:
        return 'null'
    return "'" + str(v).replace("'", "''") + "'"

def num(v):
    if v is None or v == '':
        return 0
    try:
        f = float(v)
    except (TypeError, ValueError):
        return 0
    return int(f) if f == int(f) else f

wb = openpyxl.load_workbook(XLSX, data_only=True)
ws = wb['Material']

rows = []
seen = set()
for r in range(4, ws.max_row + 1):
    code = clean(ws.cell(r, 6).value)
    if not code or code in seen:
        continue
    seen.add(code)
    cat = 'consumable' if code in CONSUMABLE else ('roll' if code in ROLL else 'general')
    rows.append({
        'code':    code,
        'desc':    clean(ws.cell(r, 7).value) or code,
        'unit':    clean(ws.cell(r, 8).value),
        'cat':     cat,
        'loc':     clean(ws.cell(r, 2).value),
        'mtype':   clean(ws.cell(r, 3).value),
        'remark':  clean(ws.cell(r, 4).value),
        'status':  clean(ws.cell(r, 5).value),
        'closing': num(ws.cell(r, 12).value),
    })

assert len(rows) == 1501, len(rows)
n_con = sum(1 for x in rows if x['cat'] == 'consumable')
n_roll = sum(1 for x in rows if x['cat'] == 'roll')
print('rows', len(rows), 'consumable', n_con, 'roll', n_roll, 'general', len(rows)-n_con-n_roll,
      file=sys.stderr)

HEAD = """-- =====================================================================
--  MMV WAREHOUSE - Danh muc vat tu day du (nguon: Material.xlsx)
--
--  {N} ma vat tu cua kho, sinh tu file Material.xlsx thang 09-2026.
--  KHONG sua file nay bang tay: sinh lai bang scripts/gen_seed_materials.py
--  khi co ban Material.xlsx moi.
--
--  Chay trong: Supabase Dashboard -> SQL Editor -> New query -> Run
--  Thu tu chay:  schema.sql  ->  seed.sql  ->  seed_materials.sql
--
--  category quyet dinh AI DUOC PHEP nhap/xuat ma do:
--    'consumable' ({NC} ma) - vat tu tieu hao, KTV tu lay o man Lay vat tu
--    'roll'       ({NR} ma) - cuon dai cat dan, o man Cat cuon
--    'general'    ({NG} ma) - moi thu con lai, CHI ADMIN duoc nhap/xuat
--                             (man Nhap/Xuat kho, /stock)
--
--  Lenh nay CHAY LAI DUOC NHIEU LAN. on conflict chi ghi de nhung cot
--  co trong Material.xlsx. description_vi, unit_price, min_stock,
--  has_expiry, expiry_date do seed.sql va nguoi dung nhap them nen
--  KHONG bi xoa khi chay lai.
-- =====================================================================

insert into materials
  (code, description, unit, category, location, mat_type, remark, stock_status,
   closing_qty, min_stock)
values
"""

TAIL = """on conflict (code) do update set
  description  = excluded.description,
  unit         = coalesce(excluded.unit, materials.unit),
  category     = excluded.category,
  location     = excluded.location,
  mat_type     = excluded.mat_type,
  remark       = excluded.remark,
  stock_status = excluded.stock_status,
  closing_qty  = excluded.closing_qty;

-- Kiem tra nhanh sau khi chay:
--   select category, count(*) from materials group by category order by 2 desc;
--   -> general {NG} / consumable {NC} / roll {NR}
"""

out = io.StringIO()
out.write(HEAD.format(N=len(rows), NC=n_con, NR=n_roll, NG=len(rows)-n_con-n_roll))
for i, x in enumerate(rows):
    # min_stock: chi vat tu tieu hao / cuon moi co nguong canh bao ton thap.
    # 'general' de 0 = khong theo doi dat hang lai (xem statusOf o Inventory.tsx),
    # neu de null thi cot lay default 20 va ca 1466 ma deu bao "Can".
    # Voi ma da co san, on conflict KHONG dung toi min_stock nen nguong
    # trong seed.sql (10-50) van giu nguyen.
    min_stock = '0' if x['cat'] == 'general' else '20'
    sep = ',' if i < len(rows) - 1 else ''
    out.write("  ({c}, {d}, {u}, {cat}, {l}, {t}, {rm}, {st}, {q}, {ms}){s}\n".format(
        c=q(x['code']), d=q(x['desc']), u=q(x['unit']), cat=q(x['cat']),
        l=q(x['loc']), t=q(x['mtype']), rm=q(x['remark']), st=q(x['status']),
        q=x['closing'], ms=min_stock, s=sep))
out.write(TAIL.format(NC=n_con, NR=n_roll, NG=len(rows)-n_con-n_roll))

with open(OUT, 'w', encoding='utf-8', newline='\n') as f:
    f.write(out.getvalue())
print('written', OUT, file=sys.stderr)
